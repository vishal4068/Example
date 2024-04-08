from docx import Document
from openpyxl import Workbook
from lxml import etree
from openpyxl.styles import Font, Alignment

# Load the Word document
docx_file = "vishal.docx"  # Replace with your file name
doc = Document(docx_file)

# Initialize Excel workbook and worksheet
workbook = Workbook()
worksheet = workbook.active

# Load the XML content from the DOCX file
xml_content = etree.fromstring(doc.part.blob)

# Extract hyperlinks
hyperlinks = xml_content.xpath('//w:hyperlink', namespaces={'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'})

worksheet.append(["Hyperlink Text", "Hyperlink URL"])

for cell in worksheet[1]:
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal='center')

hyperlink_count = 0
for hyperlink in hyperlinks:
    hyperlink_text = ""
    for element in hyperlink.iter():
        if element.tag.endswith('t') or element.tag.endswith('delText'):
            if element.text:
                hyperlink_text += element.text
            if element.tail:
                hyperlink_text += element.tail
        elif element.tag.endswith('softHyphen') or element.tag.endswith('noBreakHyphen'):
            hyperlink_text += '-'
    hyperlink_rid = hyperlink.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
    hyperlink_url = doc.part.rels[hyperlink_rid].target_ref
    worksheet.append((hyperlink_text, f'=HYPERLINK("{hyperlink_url}", "{hyperlink_text}")')) ##worksheet.append([hyperlink_text, "=HYPERLINK("{}", "{}")'.format(hyperlink_url, hyperlink_url)])
    hyperlink_count += 1

# Apply formatting to hyperlinks
for row in worksheet.iter_rows(min_row=2, min_col=2, max_col=2):
    for cell in row:
        cell.font = Font(underline='single', color='0000FF')

# Save the extracted hyperlinks to a new Excel file
output_file = "extractedHyperlinks.xlsx"
workbook.save(output_file)

print("Extracted!! Total number of Hyperlinks:", hyperlink_count)
print("Please check the current folder for the extracted Excel file.")
