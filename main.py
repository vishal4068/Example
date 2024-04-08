# from docx import Document
# from openpyxl import Workbook
# from lxml import etree

# # Load the Word document
# doc = Document('vishal.docx')

# # Create a new Excel workbook
# workbook = Workbook()
# worksheet = workbook.active

# # Parse the XML
# xml_content = etree.fromstring(doc.part.blob)

# # Find hyperlinks
# hyperlinks = xml_content.xpath('//w:hyperlink', namespaces={'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'})

# # Iterate through hyperlinks
# for hyperlink in hyperlinks:
#     hyperlink_text = hyperlink.xpath('.//w:t', namespaces={'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'})[0].text
#     hyperlink_url = hyperlink.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
#     worksheet.append([hyperlink_text, hyperlink_url])

# # Save the Excel file
# workbook.save('hyperlinks.xlsx')


from docx import Document
from openpyxl import Workbook
from lxml import etree

# Load the Word document
doc = Document('vishal.docx')

# Create a new Excel workbook
workbook = Workbook()
worksheet = workbook.active

# Parse the XML
xml_content = etree.fromstring(doc.part.blob)

# Find hyperlinks
hyperlinks = xml_content.xpath('//w:hyperlink', namespaces={'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'})

# Iterate through hyperlinks
for hyperlink in hyperlinks:
    hyperlink_text = hyperlink.xpath('.//w:t', namespaces={'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'})[0].text
    hyperlink_rid = hyperlink.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
    hyperlink_url = doc.part.rels[hyperlink_rid].target_ref
    worksheet.append([hyperlink_text, hyperlink_url])

# Save the Excel file
workbook.save('hyperlinks.xlsx')
