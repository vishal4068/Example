# from pdf2docx import Converter

# def pdf_to_docx(pdf_file, docx_file):
#     # Create a converter object
#     cv = Converter(pdf_file)

#     # Convert the PDF to Word
#     cv.convert(docx_file, start=0, end=None)

#     # Close the converter
#     cv.close()

# # Replace 'input.pdf' with the path to your PDF file
# pdf_file = 'graphics.pdf'

# # Replace 'output.docx' with the desired path for the output Word file
# docx_file = 'output.docx'

# # Convert PDF to Word
# pdf_to_docx(pdf_file, docx_file)

# print("Conversion complete.")


from docx import Document
from fpdf import FPDF

# from docx import Document

def extract_text_and_images(docx_file):
    doc = Document(docx_file)
    text_content = []
    images = []

    for paragraph in doc.paragraphs:  # Directly iterate over paragraphs
        text = ''.join([run.text for run in paragraph.runs])
        text_content.append(text)
    # For images, you'll need a different approach since images are not directly accessible as 'drawing' in paragraphs
    # You might need to explore docx's structure or use other libraries to extract images

    return text_content, images

def docx_to_pdf(docx_file, pdf_file):
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    text_content, images = extract_text_and_images(docx_file)

    pdf.add_page()
    pdf.set_font("Arial", size=12)

    for text in text_content:
        pdf.cell(200, 10, txt=text, ln=True)

    for image in images:
        if image:
            pdf.image(image, x=None, y=None, w=0, h=0)

    pdf.output(pdf_file)

docx_file = 'output.docx'
pdf_file = 'output.pdf'
docx_to_pdf(docx_file, pdf_file)

print("Conversion from DOCX to PDF complete.")